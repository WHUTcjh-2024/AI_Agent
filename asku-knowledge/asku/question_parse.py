"""Loss-aware parsing for question-first batches. Approval is a separate step."""
from __future__ import annotations

import io
import json
import re
import subprocess
import sys
import zipfile
from datetime import UTC, date, datetime
from pathlib import Path
from urllib.parse import urljoin, urlsplit

from bs4 import BeautifulSoup, NavigableString, Tag

from .fetcher import decode_bytes
from .normalizer import _extract_date, _extract_title
from .pii import detect_pii
from .question_crawl import EXTENSIONS, canonical, digest, read_jsonl, transport_url

MAIN_SELECTORS = ['#editor1', '#bfArticleContent', '.TRS_Editor', '.TRS_UEDITOR', '#vsb_content_2',
                  '#vsb_content', '.v_news_content', '.article_content', '.news-content', '.article-content',
                  '.newsDetail', '.news_detail', '.detail-content', '.content-detail', '.detail_con']
DATE_RE = re.compile(r'(20\d{2})\s*[-年/.]\s*(\d{1,2})\s*[-月/.]\s*(\d{1,2})')
COLLEGES = {'wutinfo': '信息工程学院', 'sa': '自动化学院', 'som': '管理学院', 'econ': '经济学院',
            'smse': '材料科学与工程学院', 'smee': '机电工程学院', 'sn': '航运学院',
            'stle': '交通与物流工程学院', 'naoep': '船海与能源动力工程学院', 'phymech': '物理与力学学院',
            'maths': '数学与统计学院', 'ssci': '理学院', 'sree': '资源与环境工程学院',
            'csai': '计算机与人工智能学院', 'auto': '汽车工程学院', 'wenfa': '法学与人文社会学院',
            'sfl': '外国语学院', 'sen': '创业学院', 'sie': '国际教育学院', 'sccels': '化学化工与生命科学学院',
            'ismse': '材料科学与工程国际化示范学院', 'amucwut': '艾克斯马赛学院', 'sports': '体育学院'}
DEPARTMENTS = {'jwc': '本科生院', 'nic': '网络信息中心', 'stuplaza': '学生工作部',
               'dzb': '党政办公室', 'lib': '图书馆', 'zs': '本科生招生办公室',
               'xxgk': '信息公开网', 'www': '武汉理工大学', 'youth': '校团委', 'jcc': '财务处',
               'gd': '研究生院', 'hp': '校医院'}
RESULT_RE = re.compile(r'名单|录取结果|转专业结果|评选结果|拟录取|获奖.{0,4}公示|等\d+名学生')
FORM_RE = re.compile(r'申请表|审批表|登记表|汇总表|承诺书|申请单|模板')


def tidy(text: str) -> str:
    text = text.replace('\r\n', '\n').replace('\r', '\n').replace('\xa0', ' ').replace('\u3000', ' ')
    text = re.sub(r'[\u200b\ufeff\x00\x01\x07]', '', text)
    text = re.sub(r'[ \t]+', ' ', text)
    return re.sub(r'\n{3,}', '\n\n', '\n'.join(line.strip() for line in text.splitlines())).strip()


def table_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ''
    width = max(map(len, rows))
    if not width:
        return ''
    rows = [[tidy(str(v)).replace('\n', '<br>').replace('|', '\\|') for v in row] + ['']*(width-len(row)) for row in rows]
    return '\n'.join(['| '+' | '.join(rows[0])+' |', '| '+' | '.join(['---']*width)+' |']+
                     ['| '+' | '.join(r)+' |' for r in rows[1:]])


def html_table(table: Tag) -> str:
    """Expand row/col spans so a reader never loses the scope of a merged cell."""
    grid: dict[tuple[int, int], str] = {}
    max_row = max_col = 0
    for i, row in enumerate(table.find_all('tr')):
        j = 0
        for cell in row.find_all(['td', 'th'], recursive=False):
            while (i, j) in grid:
                j += 1
            text = tidy(cell.get_text(' ', strip=True))
            rs = min(int(cell.get('rowspan', 1) or 1), 200)
            cs = min(int(cell.get('colspan', 1) or 1), 60)
            for y in range(i, i+rs):
                for x in range(j, j+cs):
                    grid[y, x] = text
                    max_row, max_col = max(max_row, y+1), max(max_col, x+1)
            j += cs
    return table_markdown([[grid.get((i, j), '') for j in range(max_col)] for i in range(max_row)])


def html_markdown(node: Tag, base: str) -> str:
    def render(n):
        if isinstance(n, NavigableString):
            return str(n)
        if not isinstance(n, Tag) or n.name in {'script', 'style', 'noscript', 'iframe', 'form', 'input', 'button'}:
            return ''
        if n.name == 'table':
            return '\n\n'+html_table(n)+'\n\n'
        if n.name == 'img':
            return '\n\n[原文图片：'+urljoin(base, n.get('src', ''))+']\n\n'
        if n.name == 'br':
            return '\n'
        inner = ''.join(render(child) for child in n.children)
        if n.name == 'a' and n.get('href'):
            href = urljoin(base, n['href'])
            return '['+tidy(inner)+']('+href+')' if href.startswith(('http:', 'https:')) and tidy(inner) else inner
        if n.name in {'h1', 'h2', 'h3', 'h4', 'h5', 'h6'}:
            return '\n\n'+'#'*min(int(n.name[1])+1, 6)+' '+inner+'\n\n'
        if n.name == 'li':
            return '\n- '+inner+'\n'
        if n.name in {'p', 'div', 'section', 'article', 'blockquote', 'ul', 'ol'}:
            return '\n\n'+inner+'\n\n'
        return inner
    return tidy(render(node))


def normalize_title(title: str) -> str:
    title = tidy(title)
    title = re.sub(r'^20\d\d-\d\d-\d\d\s+', '', title)
    title = re.split(r'\s*(?:发布日期|发布时间|更新时间)[:：]', title)[0]
    title = re.sub(r'[-_—]\s*武汉理工大学[^《》]{0,40}$', '', title)
    return title.strip()


def parse_html(data: bytes, url: str, hint: str = '') -> dict:
    html, encoding = decode_bytes(data)
    soup = BeautifulSoup(html, 'lxml')
    title = normalize_title(_extract_title(soup))
    if not title or title.endswith(('欢迎您', '欢迎您！', '学工部（处） 武装部')) or title in {
            '武汉理工大学本科生院', '武汉理工大学图书馆', '武汉理工大学党委学工部（处） 武装部'}:
        title = normalize_title(hint)
    container = None
    selector = None
    for sel in MAIN_SELECTORS:
        node = soup.select_one(sel)
        if node is not None and (node.get_text(strip=True) or node.select('img')):
            container, selector = node, sel
            break
    flags = []
    if container is None:
        # Keep only explicitly named detail bodies; never fall back to whole body.
        for sel in ['.text', '.nr', '.con', '.detail', '.content']:
            for node in soup.select(sel):
                if len(node.get_text(strip=True)) > 150 and not node.select('nav') and len(node.select('a')) < 12:
                    container, selector = node, sel
                    flags.append('fallback_content_selector')
                    break
            if container is not None:
                break
    published = _extract_date(soup, soup.get_text('\n', strip=True))
    date_evidence = 'page_publish_field' if published else None
    if not published:
        for meta in soup.select('meta[content]'):
            if (meta.get('name') or meta.get('property') or '').lower() in {'articlepubdate', 'pubdate', 'publishdate', 'date', 'article:published_time'}:
                m = DATE_RE.search(meta['content'])
                if m:
                    try:
                        published = date(*map(int, m.groups())).isoformat()
                        date_evidence = 'meta:'+str(meta.get('name') or meta.get('property'))
                    except ValueError:
                        pass
    if not published:
        flags.append('publication_date_unknown')
    if container is None:
        flags.append('no_article_body')
        text = ''
    else:
        text = html_markdown(container, url)
    links = []
    for a in soup.select('a[href]'):
        href = canonical(transport_url(urljoin(url, a['href'])))
        if Path(urlsplit(href).path).suffix.lower() in EXTENSIONS:
            links.append({'url': href, 'filename': tidy(a.get_text(' ', strip=True))})
    if '[原文图片：' in text:
        flags.append('image_evidence_required')
    return {'title': title, 'clean_content': text, 'published_at': published,
            'publish_date_evidence': date_evidence, 'quality_flags': flags,
            'attachment_links': list({x['url']:x for x in links}.values()), 'parser': 'html:'+str(selector), 'encoding': encoding}


def parse_docx(data: bytes) -> tuple[str, list[str]]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(io.BytesIO(data))
    blocks = []
    for element in doc.element.body:
        if element.tag.endswith('}p'):
            p = Paragraph(element, doc)
            blocks.append(p.text)
        elif element.tag.endswith('}tbl'):
            table = Table(element, doc)
            blocks.append(table_markdown([[c.text for c in row.cells] for row in table.rows]))
    flags = ['embedded_images'] if doc.inline_shapes else []
    return tidy('\n\n'.join(blocks)), flags


def parse_file(path: Path, cache: Path) -> tuple[str, list[str], dict]:
    suffix = path.suffix.lower()
    data = path.read_bytes()
    if data.startswith(b'%PDF'):
        import pymupdf
        with pymupdf.open(stream=data, filetype='pdf') as doc:
            blocks, blank_pages, tables = [], [], 0
            for idx, page in enumerate(doc):
                text = page.get_text(sort=True)
                if len(re.sub(r'\s', '', text)) < 20:
                    blank_pages.append(idx+1)
                # Keep full text and preserve tabular structures on table pages.
                found = page.find_tables().tables
                if found:
                    table_boxes = [t.bbox for t in found]
                    others = []
                    for block in page.get_text('blocks', sort=True):
                        if not any(pymupdf.Rect(block[:4]).intersects(pymupdf.Rect(box)) for box in table_boxes):
                            others.append((block[1], block[4]))
                    others.extend((t.bbox[1], table_markdown(t.extract())) for t in found)
                    text = '\n\n'.join(x[1] for x in sorted(others, key=lambda x:x[0]))
                    tables += len(found)
                blocks.append('### 第 '+str(idx+1)+' 页\n\n'+text)
            flags = (['pdf_ocr_required'] if blank_pages else []) + (['pdf_table_layout_review'] if tables else [])
            return tidy('\n\n'.join(blocks)), flags, {'pages':len(doc), 'blank_pages':blank_pages, 'tables':tables}
    if suffix == '.docx':
        text, flags = parse_docx(data)
        return text, flags, {}
    if suffix == '.doc':
        converted = cache / (digest(data)+'.docx')
        if not converted.exists():
            script = "import sys,win32com.client;w=win32com.client.DispatchEx('Word.Application');w.Visible=False;w.DisplayAlerts=0;w.AutomationSecurity=3\ntry:\n d=w.Documents.Open(sys.argv[1],ReadOnly=True,AddToRecentFiles=False);d.SaveAs2(sys.argv[2],FileFormat=16);d.Close(False)\nfinally:w.Quit()"
            subprocess.run([sys.executable, '-c', script, str(path.resolve()), str(converted.resolve())], timeout=60, check=True, capture_output=True)
        text, flags = parse_docx(converted.read_bytes())
        return text, flags, {'converted_from':'doc'}
    if suffix in {'.xlsx', '.xls'}:
        sheets = []
        if suffix == '.xlsx':
            import openpyxl
            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                for sheet in book:
                    rows = [[str(v) if v is not None else '' for v in row] for row in sheet.iter_rows(values_only=True)]
                    sheets.append((sheet.title, rows))
            finally:
                book.close()
        else:
            import xlrd
            book = xlrd.open_workbook(file_contents=data)
            sheets = [(s.name, [[str(v) if v is not None else '' for v in s.row_values(i)] for i in range(s.nrows)]) for s in book.sheets()]
        text = '\n\n'.join('## '+name+'\n\n'+table_markdown([r for r in rows if any(r)]) for name, rows in sheets)
        return tidy(text), [], {'sheets':[n for n,_ in sheets]}
    if suffix in {'.txt', '.csv'}:
        return tidy(decode_bytes(data)[0]), [], {}
    return '', ['unsupported_or_image'], {}


def archive_members(path: Path, cache: Path) -> list[tuple[str, Path]]:
    """Extract only bounded document bytes to hashed paths; never trust ZIP paths."""
    result = []
    with zipfile.ZipFile(path) as zf:
        infos = zf.infolist()
        if len(infos)>200 or sum(i.file_size for i in infos)>60*1024*1024:
            raise ValueError('archive_budget_exceeded')
        for info in infos:
            name = info.filename
            if not info.flag_bits & 0x800:
                try:
                    name = name.encode('cp437').decode('gbk')
                except (UnicodeError, LookupError):
                    pass
            p = Path(name.replace('\\','/'))
            if info.is_dir() or p.suffix.lower() not in EXTENSIONS- {'.zip'}:
                continue
            if p.is_absolute() or '..' in p.parts or re.match(r'^[A-Za-z]:', name) or ((info.external_attr>>16)&0o170000)==0o120000:
                raise ValueError('unsafe_archive_member')
            if info.file_size>30*1024*1024 or info.file_size/max(1,info.compress_size)>250:
                raise ValueError('archive_member_budget')
            data = zf.read(info)
            dest = cache/(digest(data)+p.suffix.lower())
            if not dest.exists():
                dest.write_bytes(data)
            result.append((name,dest))
    return result


def source_metadata(url: str, title: str) -> dict:
    host = (urlsplit(url).hostname or '').split('.')[0]
    if host in COLLEGES:
        return {'source_name': COLLEGES[host], 'department': COLLEGES[host], 'college': COLLEGES[host],
                'scope':'college', 'authority':'OFFICIAL_COLLEGE'}
    if host == 'i':
        department = '本科生院' if '/jwc/' in url else ('图书馆' if '/tsg/' in url else ('学工部' if '/xgb/' in url else None))
        bracket = re.match(r'【([^】]+)】', title)
        if not department and bracket:
            department = bracket[1]
        return {'source_name':'武汉理工大学综合信息网', 'department':department, 'college':None,
                'scope':'school' if department in {'本科生院', '图书馆', '学工部', '财务处', '党政办'} else None,
                'authority':'OFFICIAL_DEPARTMENT'}
    name = DEPARTMENTS.get(host)
    return {'source_name':name, 'department':name, 'college':None, 'scope':'school' if name else None,
            'authority':'OFFICIAL_SCHOOL' if host in {'xxgk','www','dzb'} else 'OFFICIAL_DEPARTMENT'}


def enrich(row: dict, patterns: dict) -> dict:
    title, text = row['title'], row['clean_content']
    topics = [k for k,p in patterns.items() if p.search(title)]
    if not topics and row.get('parent_topic'):
        topics = row['parent_topic']
    primary_rules = [
        ('transfer_major', r'转专业'),
        ('course_selection', r'选课|课程预选|补退选'),
        ('cet', r'四[、，,]?六级|CET|大学英语.{0,5}级'),
        ('makeup_retake', r'补考|缓考|补（缓）考|重修|重考|免听'),
        ('student_status', r'学籍管理|休学|复学|退学|转学|延长学习|提前毕业|毕业资格|学士学位'),
        ('comprehensive_evaluation', r'综合素质测评|综合测评|综测'),
        ('scholarship', r'奖学金|奖（助）学金|奖助学金|学生奖励|评先'),
        ('library', r'图书馆|借阅|续借|入馆|座位|研讨室|校外访问'),
        ('academic_calendar', r'校历|放假|寒假|暑假|开学'),
        ('campus_card', r'校园卡|一卡通'),
        ('exam', r'考试|考核|成绩'),
    ]
    primary = next((topic for topic, pattern in primary_rules if re.search(pattern, title, re.IGNORECASE)), None)
    if primary:
        topics = [primary] + [topic for topic in topics if topic != primary]
    row['topics'] = topics
    row['topic'] = topics[0] if topics else 'other'
    row['subtopics'] = [k for k, p in {'makeup_exam':r'补考|补（缓）考', 'deferred_exam':r'缓考|补（缓）考',
                                    'retake':r'重修', 'reexam':r'重考'}.items() if re.search(p, title+'\n'+text)]
    row.update(source_metadata(row.get('parent_page_url') or row['official_url'], title))
    row['education_level'] = 'BOTH' if row['topic'] in {'library','campus_card','academic_calendar'} else 'UNDERGRADUATE'
    if re.search(r'研究生|硕士|博士', title):
        row['education_level'] = 'GRADUATE'
    if '留学生' in title or '国际学生' in title or row.get('college')=='国际教育学院':
        row['quality_flags'].append('special_audience_requires_review')
    row['audience'] = row['education_level']
    dtype = 'notice'
    if re.search('办法|规定|细则|守则', title):
        dtype = 'regulation'
    elif re.search('指南|说明|操作|手册|流程|服务|开放时间', title):
        dtype = 'guide'
    if '校历' in title:
        dtype = 'calendar'
    if row.get('is_attachment'):
        dtype = 'attachment' if dtype == 'notice' else dtype
    if FORM_RE.search(title):
        dtype = 'form'
    if RESULT_RE.search(title):
        dtype = 'result'
    row['document_type'] = dtype
    row['source_rank'] = 'S' if dtype=='regulation' else ('B' if row['scope']=='college' else 'A')
    pii = detect_pii(text, title=title, tables_markdown=text)
    row['contains_pii'] = bool(pii.categories) or dtype == 'result'
    row['pii_categories'] = pii.categories + (['result_roster_risk'] if dtype=='result' else [])
    row['public_contacts_present'] = bool(re.search(r'(?<!\d)1[3-9]\d{9}(?!\d)|[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', text))
    if row['contains_pii']:
        row['quality_flags'].append('pii_review')
    if re.search(r'\|[^\n]*(姓名|学号)[^\n]*\|', text) and re.search(r'(?<!\d)\d{10,18}(?!\d)', text):
        row['contains_pii'] = True
        row['quality_flags'].append('roster_table')
    if row['public_contacts_present']:
        row['quality_flags'].append('contact_context_review')
    if len(re.sub(r'\s', '', text)) < 80:
        row['quality_flags'].append('short_content')
    if not row['scope'] or not row['department']:
        row['quality_flags'].append('source_scope_unknown')
    if '\ufffd' in text or '\ufffd' in title:
        row['quality_flags'].append('unicode_corruption')
    if row['education_level']=='GRADUATE':
        row['quality_flags'].append('outside_undergraduate_p0')
    years = re.search(r'(20\d\d)\s*[-—–至]\s*(20\d\d)学年', title)
    year = re.search(r'(20\d\d)年', title)
    row['academic_year'] = '-'.join(years.groups()) if years else (year[1] if year else None)
    row['semester'] = '1' if '第一学期' in title or '秋季' in title else ('2' if '第二学期' in title or '春季' in title else None)
    row['effective_from'] = None
    row['effective_to'] = None
    row['version_status'] = 'unknown'
    if row['academic_year'] and dtype in {'notice','calendar','result'} and int(row['academic_year'].split('-')[-1]) < 2026:
        row['version_status'] = 'historical'
    row['version_group'] = digest('|'.join([str(row['scope']),str(row['college']),row['topic'],
                                           re.sub(r'20\d\d|\d+|\s', '', title),dtype]))[:24]
    row['supersedes'] = None
    row['superseded_by'] = None
    row['content_hash'] = digest(tidy(text))
    row['question_tags'] = []
    row['questions'] = []
    row['review_status'] = 'REVIEW'
    row['rag_eligible'] = False
    row['cleaned_at'] = datetime.now(UTC).isoformat()
    row['quality_flags'] = sorted(set(row['quality_flags']))
    return row


def build_candidates(crawl: Path, config: Path, inventory: Path, output: Path, rendered: Path | None = None):
    import yaml
    cfg = yaml.safe_load(config.read_text(encoding='utf-8'))
    patterns = {t['id']:re.compile(t['title_pattern'], re.IGNORECASE) for t in cfg['topics']}
    output.mkdir(parents=True, exist_ok=True)
    cache = output/'parse-cache'
    cache.mkdir(exist_ok=True)
    def url_key(u):
        return canonical(u).split('://',1)[-1]
    old_ids = {url_key(r['source_url']):r['id'] for r in read_jsonl(inventory)}
    frontier = {r['url']:r for r in read_jsonl(crawl/'frontier.jsonl')}
    ledger = {r['requested_url']:r for r in read_jsonl(crawl/'fetch-ledger.jsonl')}
    rows = []
    for url, f in frontier.items():
        fetched = ledger.get(url)
        if not fetched or not fetched['ok']:
            continue
        path = crawl/fetched['raw_file']
        row = {'asku_document_id':old_ids.get(url_key(url)) or 'whq_'+digest(url)[:28], 'school_id':'whut',
               'official_url':fetched['url'], 'canonical_url':url, 'raw_file':str(path.resolve()),
               'raw_sha256':fetched['sha256'], 'crawl_at':fetched['crawl_at'], 'source_status':'available',
               'mime_type':fetched['mime_type'], 'filename':f['title_hint'] or Path(urlsplit(url).path).name,
               'attachment_url':url if f['kind']!='page' else None, 'parent_page_url':f['parent_urls'][0] if f['parent_urls'] else None,
               'parent_page_urls':f['parent_urls'], 'is_attachment':f['kind']!='page', 'quality_flags':[],
               'knowledge_bundle_id':'kbq_'+digest(f['parent_urls'][0] if f['parent_urls'] else url)[:24]}
        try:
            if f['kind']=='page':
                row.update(parse_html(path.read_bytes(), fetched['url'], f['title_hint']))
            else:
                key = digest(path.read_bytes())
                cached = cache/(key+'.json')
                if RESULT_RE.search(f['title_hint']):
                    result = {'clean_content':'', 'quality_flags':['roster_not_parsed'], 'parse_details':{'retained_raw_only':True}}
                elif cached.exists():
                    result = json.loads(cached.read_text(encoding='utf-8'))
                else:
                    text, flags, details = parse_file(path, cache)
                    result = {'clean_content':text, 'quality_flags':flags, 'parse_details':details}
                    cached.write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8')
                row.update(result)
                row.update(title=normalize_title(f['title_hint'] or path.name), published_at=None,
                           publish_date_evidence=None, parser='attachment:'+path.suffix.lower())
        except Exception as exc:  # noqa: BLE001 - each source artifact must fail closed, independently.
            row.update(title=f['title_hint'], clean_content='', published_at=None, parser='failed',
                       parse_error=type(exc).__name__, quality_flags=['parse_failure'])
        rows.append(row)
    # ZIP containers stay in audit; their policy documents retain original member provenance.
    children = []
    for row in rows:
        path = Path(row['raw_file'])
        if path.suffix.lower() != '.zip' or not re.search('转专业|测评|奖学金|选课|校历|学籍',row['title']):
            continue
        try:
            members = archive_members(path, cache)
            row['archive_members']=[n for n,_ in members]
            for name, member in members:
                child = dict(row)
                child['asku_document_id']='whq_'+digest(row['canonical_url']+'#'+name)[:28]
                child['archive_member']=name
                child['archive_url']=row['official_url']
                child['canonical_url']=row['canonical_url']+'#member='+digest(name)[:16]
                child['filename']=name
                child['title']=normalize_title(Path(name).name)
                child['raw_file']=str(member.resolve())
                child['raw_sha256']=digest(member.read_bytes())
                child.pop('archive_members',None)
                cached=cache/(child['raw_sha256']+'.json')
                if cached.exists():
                    parsed=json.loads(cached.read_text(encoding='utf-8'))
                else:
                    text, flags, details=parse_file(member,cache)
                    parsed={'clean_content':text,'quality_flags':flags,'parse_details':details}
                    cached.write_text(json.dumps(parsed,ensure_ascii=False),encoding='utf-8')
                child.update(parsed)
                child['parser']='archive_member:'+member.suffix.lower()
                children.append(child)
        except Exception as exc:  # noqa: BLE001 - a malformed archive must not stop the batch.
            row['quality_flags'].append('archive_'+type(exc).__name__)
    rows.extend(children)
    # Browser DOM capture is a source artifact, not synthesized policy text.
    if rendered:
        for p in sorted(rendered.glob('library-rendered-*.json')):
            content = p.read_text(encoding='utf-8')
            try:
                r = json.loads(content)
            except json.JSONDecodeError:
                r = json.loads(content.split('### Result\n')[1].split('\n### Ran')[0])
            url = canonical(r['url'])
            body = '<html><head><title>'+r['title']+'</title></head><body>'+r['html']+'</body></html>'
            raw = cache/(digest(body)+'.html')
            raw.write_text(body, encoding='utf-8')
            row = {
                'asku_document_id': 'whq_'+digest(url)[:28], 'school_id': 'whut',
                'official_url': url, 'canonical_url': url, 'raw_file': str(raw.resolve()),
                'raw_sha256': digest(body), 'crawl_at': r['observed_at'], 'source_status': 'available',
                'mime_type': 'text/html', 'filename': r['title'], 'attachment_url': None,
                'parent_page_url': None, 'parent_page_urls': [], 'is_attachment': False,
                'knowledge_bundle_id': 'kbq_'+digest(url)[:24],
            }
            row.update(parse_html(body.encode(),url,r['title']))
            row['parser']='browser_dom:#editor1'
            rows.append(row)
    parents = {canonical(transport_url(r['official_url'])):r for r in rows if not r['is_attachment']}
    for row in rows:
        if row.get('parent_page_url'):
            p = parents.get(canonical(transport_url(row['parent_page_url'])))
            if p:
                row['published_at']=p['published_at']
                row['publish_date_evidence']='parent_page:'+p['asku_document_id']
                row['parent_topic']=[k for k,pat in patterns.items() if pat.search(p['title'])]
                row['parent_title']=p['title']
        enrich(row,patterns)
    destination = output/'candidates.jsonl'
    destination.write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in rows),encoding='utf-8')
    from collections import Counter
    print(json.dumps({'parsed':len(rows),'topics':Counter(r['topic'] for r in rows),
                      'flags':Counter(f for r in rows for f in r['quality_flags'])},ensure_ascii=False),flush=True)
    return rows
